import json

import pytest

from conftest import BLOCKS, _block
from ocr import nb, ops, prompts, store


# ── notebook loader ───────────────────────────────────────────────────────
def test_every_role_matches_one_cell():
    for role in nb.ROLES:
        assert nb.cell(role)


def test_input_patch_touches_only_two_lines():
    raw, patched = nb.cell("config"), nb.source("config")
    diff = [(a, b) for a, b in zip(raw.split("\n"), patched.split("\n")) if a != b]
    assert len(diff) == 2 and diff[0][1].startswith("PDF_PATH") and diff[1][1].startswith("PAGES")


def test_run_script_has_snapshot_before_final_cell():
    s = nb.run_script()
    compile(s, "run", "exec")
    assert s.index("blocks_pre.json") < s.index("ordered_blocks = fix_blocks(ordered_blocks)")


def test_changed_notebook_fails_loudly(tmp_path):
    book = json.loads(nb.NOTEBOOK.read_text(encoding="utf-8"))
    book["cells"] = [c for c in book["cells"] if "def run_single_gemini" not in "".join(c["source"])]
    p = tmp_path / "x.ipynb"
    p.write_text(json.dumps(book), encoding="utf-8")
    with pytest.raises(nb.NotebookError, match="single"):
        nb.cell("single", p)


# ── history ───────────────────────────────────────────────────────────────
def test_events_undo_redo(job):
    jid = job.name
    store.push(jid, {"op": "set_fields", "block": "p1_b1", "fields": {"formatted_text": "<P>X</P>"}})
    assert store.current_blocks(jid)[0]["formatted_text"] == "<P>X</P>"
    store.move_cursor(jid, -1)
    assert store.current_blocks(jid)[0]["formatted_text"].startswith('<P align="center"')
    store.move_cursor(jid, +1)
    assert store.current_blocks(jid)[0]["formatted_text"] == "<P>X</P>"
    store.move_cursor(jid, -1)
    store.push(jid, {"op": "set_fields", "block": "p1_b1", "fields": {"formatted_text": "<P>Y</P>"}})   # drops the undone edit
    assert len(store.history(jid)["events"]) == 1


def test_reorder_rules(job):
    jid = job.name
    ids = [b["block_id"] for b in BLOCKS]
    store.push(jid, {"op": "reorder", "order": ["p1_b2", "p1_b1"] + ids[2:]})            # within page 1: fine
    assert [b["block_id"] for b in store.current_blocks(jid)][:2] == ["p1_b2", "p1_b1"]
    with pytest.raises(ValueError, match="within their own page"):
        store.push(jid, {"op": "reorder", "order": ids[1:] + ids[:1]})                     # moves a page-1 block to the end
    with pytest.raises(ValueError, match="every block"):
        store.push(jid, {"op": "reorder", "order": ids[:-1]})


def test_edit_validation(job):
    jid = job.name
    for bad in ({"op": "set_fields", "block": "nope", "fields": {}},
                {"op": "set_fields", "block": "p1_b1", "fields": {"page": 9}},
                {"op": "add_block", "block": {"block_id": "p1_b1", "page": 1, "formatted_text": ""}, "index": 0},
                {"op": "bogus"}):
        with pytest.raises(ValueError):
            store.push(jid, bad)


# ── rebuild ───────────────────────────────────────────────────────────────
def test_rebuild_joins_page_break_and_reports_it(job):
    ops.rebuild(job)
    ui = store.read_json(job / "ui.json")
    from docx import Document
    text = "\n".join(p.text for p in Document(job / "output.docx").paragraphs)
    assert "proportional to the potential difference across its ends." in text
    assert ui["blocks"]["p2_b1"]["absorbed"] and ui["blocks"]["p2_b1"]["absorbed_by"] == "p1_b4"
    assert not ui["blocks"]["p1_b4"]["absorbed"]
    assert ui["blocks"]["p1_b4"]["markup"] == BLOCKS[3]["formatted_text"]          # the screen edits pre-fix text
    assert (job / "pages" / "1.jpg").exists() and (job / "pages" / "2.jpg").exists()
    assert store.status(job.name)["built_rev"] == 0


def test_rebuild_applies_edits_and_order(job):
    jid = job.name
    store.push(jid, {"op": "set_fields", "block": "p2_b2", "fields": {"formatted_text": "<P bold=\"true\">Changed.</P>"}})
    ops.rebuild(job)
    from docx import Document
    assert "Changed." in "\n".join(p.text for p in Document(job / "output.docx").paragraphs)
    assert store.status(jid)["built_rev"] == 1


# ── review ────────────────────────────────────────────────────────────────
QWEN = "qwen/qwen3-vl-30b-a3b-instruct"          # the model the notebook names for its 'classify' stage


def test_fix_formatting_asks_qwen_with_the_instruction_first(job, fake_models):
    res = ops.review(job, {"mode": "format", "blocks": ["p2_b2"], "note": "make it bold"})
    assert res["blocks"][0]["after"] == "<P>Next paragraph.</P>"
    assert res["blocks"][0]["fields"] == {"formatted_text": "<P>Next paragraph.</P>"}
    body = fake_models[0][1]
    sent = body["messages"][0]["content"][0]["text"]
    assert body["model"] == QWEN and res["model"] == QWEN                          # no Gemini anywhere in review
    assert sent.index("THE PERSON'S INSTRUCTION") < sent.index("make it bold") < sent.index("TAG REFERENCE") < sent.index("INPUT:")
    assert "(A) Words" in sent and "(C) Layout" in sent                             # the labels the notebook's closing sentence points at
    assert '<BLOCK id="p2_b2" type="text"' in sent and "Next paragraph." in sent
    assert any(c.get("type") == "image_url" for c in body["messages"][0]["content"])    # page image attached


def test_without_a_note_the_default_instruction_is_used(job, fake_models):
    ops.review(job, {"mode": "format", "blocks": ["p2_b2"]})
    assert prompts.load("default_instruction.txt").strip() in fake_models[0][1]["messages"][0]["content"][0]["text"]


def test_the_model_sees_each_blocks_current_tags_and_the_answer_is_cleaned(job, fake_models):
    store.push(job.name, {"op": "set_fields", "block": "p2_b2", "fields": {"formatted_text": '<P align="center">Next paragraph.</P>'}})
    res = ops.review(job, {"mode": "format", "blocks": ["p2_b2"], "note": "make it bold"})
    sent = fake_models[0][1]["messages"][0]["content"][0]["text"]
    assert '<CURRENT_TAGS><P align="center">Next paragraph.</P></CURRENT_TAGS>' in sent       # the model starts from the centred paragraph
    assert "CURRENT_TAGS" not in res["blocks"][0]["after"]


def test_fix_formatting_refuses_pictures_and_unknown_blocks(job, fake_models):
    with pytest.raises(ValueError, match="unknown"):
        ops.review(job, {"mode": "format", "blocks": ["nope"]})
    with pytest.raises(ValueError, match="not allowed"):
        ops.review(job, {"mode": "format", "blocks": ["p2_b2"], "model": "evil/model"})


def test_reread_reads_the_crop_then_formats(job, fake_models):
    res = ops.review(job, {"mode": "reread", "blocks": ["p2_b2"], "note": "this line is Telugu"})
    b = res["blocks"][0]
    assert b["after"] == "<P>alpha beta</P>" and b["fields"]["is_picture"] is False
    urls = [u for u, _ in fake_models]
    assert any("bodhan" in u for u in urls) and any("openrouter" in u for u in urls)
    last = fake_models[-1][1]
    assert last["model"] == QWEN
    sent = last["messages"][0]["content"][0]["text"]
    assert "REVIEWER NOTE" in sent and "this line is Telugu" in sent               # re-read keeps the notebook's own prompt, plus the hint


def test_the_backup_ocr_also_uses_qwen_not_gemini(job):
    ns = ops.nb.load(lift=("indic_ocr", "indic_fmt"))
    ops._configure(ns, "", None, "reread")
    assert ns["OCR_FALLBACK_MODEL"] == QWEN and ns["COMBINED_MODELS"] == [QWEN]


def test_a_model_outside_the_notebooks_list_is_refused(job, fake_models):
    with pytest.raises(ValueError, match="not allowed"):
        ops.review(job, {"mode": "format", "blocks": ["p2_b2"], "model": "evil/model"})


def test_newbox_proposes_block_in_reading_position(job, fake_models):
    res = ops.newbox(job, {"page": 1, "bbox_norm": [.1, .60, .9, .70]})
    assert res["new_block"]["block_id"] == "p1_u1" and res["new_block"]["formatted_text"] == "<P>alpha beta</P>"
    assert res["index"] == 3                                  # after the table (starts .35), before p1_b4 (starts .85)
    jid = job.name
    store.push(jid, {"op": "add_block", "block": res["new_block"], "index": res["index"]})
    assert [b["block_id"] for b in store.current_blocks(jid)][2:5] == ["p1_b3", "p1_u1", "p1_b4"]
    with pytest.raises(ValueError, match="inside the page"):
        ops.newbox(job, {"page": 1, "bbox_norm": [.5, .5, .4, .6]})


# ── review: one unusable block must not sink the others ───────────────────
def _add(jid, block, index=0):
    store.push(jid, {"op": "add_block", "block": block, "index": index})


def test_a_block_that_cannot_be_formatted_is_skipped_not_fatal(job, fake_models):
    _add(job.name, _block("p1_empty", 1, [.1, .6, .9, .7], "text", ""))
    res = ops.review(job, {"mode": "format", "blocks": ["p1_empty", "p2_b2"]})
    assert [b["id"] for b in res["blocks"]] == ["p2_b2"]
    assert res["skipped"] and res["skipped"][0]["id"] == "p1_empty" and "Re-read" in res["skipped"][0]["reason"]


def test_if_nothing_can_be_done_the_reason_is_the_error(job, fake_models):
    _add(job.name, _block("p1_empty", 1, [.1, .6, .9, .7], "text", ""))
    with pytest.raises(ValueError, match="p1_empty: nothing to format"):
        ops.review(job, {"mode": "format", "blocks": ["p1_empty"]})


def test_a_table_goes_to_the_model_as_its_current_table_tags(job, fake_models):
    tbl = "<TABLE><ROW><CELL>Quantity</CELL><CELL>Unit</CELL></ROW><ROW><CELL>Current</CELL><CELL>Ampere</CELL></ROW></TABLE>"
    _add(job.name, _block("p1_tbl", 1, [.1, .6, .9, .7], "table", tbl))
    ops.review(job, {"mode": "format", "blocks": ["p1_tbl"]})
    sent = fake_models[0][1]["messages"][0]["content"][0]["text"]
    assert f"<CURRENT_TAGS>{tbl}</CURRENT_TAGS>" in sent and "Quantity Unit Current Ampere" in sent, sent[sent.index("INPUT:"):][:700]


def test_many_blocks_go_in_parallel_chunks_and_all_come_back(job, fake_models):
    for k in range(17):
        _add(job.name, _block(f"p1_x{k}", 1, [.1, .02 * k, .9, .02 * k + .015], "text", f"<P>Block number {k}.</P>"), index=k)
    res = ops.review(job, {"mode": "format", "blocks": [f"p1_x{k}" for k in range(17)]})
    assert len(res["blocks"]) == 17 and not res["skipped"]
    assert len([u for u, _ in fake_models if "openrouter" in u]) == 3     # 8 + 8 + 1


def test_missing_key_is_said_plainly(job, fake_models, monkeypatch):
    monkeypatch.delenv("OPENROUTER_API_KEY")
    with pytest.raises(ValueError, match="no OPENROUTER_API_KEY"):
        ops.review(job, {"mode": "format", "blocks": ["p2_b2"]})


# ── delete ────────────────────────────────────────────────────────────────
def test_delete_removes_blocks_and_undo_brings_them_back(job):
    jid = job.name
    ops.rebuild(job)
    store.push(jid, {"op": "delete", "blocks": ["p1_b2", "p2_b2"]})
    assert [b["block_id"] for b in store.current_blocks(jid)] == ["p1_b1", "p1_b3", "p1_b4", "p2_b1"]
    live = store.light_doc(jid)
    assert "p1_b2" not in live["blocks"] and live["order"] == ["p1_b1", "p1_b3", "p1_b4", "p2_b1"] and live["dirty"] is True
    store.move_cursor(jid, -1)
    back = store.light_doc(jid)
    assert "p1_b2" in back["blocks"] and back["blocks"]["p1_b2"]["bbox_norm"] == [.1, .15, .9, .30]     # position known without a rebuild


def test_delete_is_checked(job):
    with pytest.raises(ValueError, match="unknown"):
        store.push(job.name, {"op": "delete", "blocks": ["nope"]})
    with pytest.raises(ValueError, match="at least one block"):
        store.push(job.name, {"op": "delete", "blocks": [b["block_id"] for b in BLOCKS]})


def test_a_deleted_block_leaves_the_word_file(job):
    ops.rebuild(job)
    store.push(job.name, {"op": "delete", "blocks": ["p1_b1"]})
    ops.rebuild(job)
    from docx import Document
    assert "Model Question Paper" not in "\n".join(p.text for p in Document(job / "output.docx").paragraphs)


def test_the_instant_view_needs_no_notebook_and_tracks_edits(job):
    ops.rebuild(job)
    store.push(job.name, {"op": "set_fields", "block": "p1_b1", "fields": {"formatted_text": "<P>Renamed</P>"}})
    store.push(job.name, {"op": "reorder", "order": ["p1_b2", "p1_b1", "p1_b3", "p1_b4", "p2_b1", "p2_b2"]})
    v = store.light_doc(job.name)
    assert v["blocks"]["p1_b1"]["markup"] == "<P>Renamed</P>" and v["order"][:2] == ["p1_b2", "p1_b1"] and v["dirty"] is True
    assert v["blocks"]["p2_b1"]["absorbed_by"] == "p1_b4"          # untouched flags carry over from the last rebuild


# ── the prompts are files a person can edit ──────────────────────────────
def test_prompt_files_are_created_and_edits_are_used_at_once(job, fake_models):
    ops.review(job, {"mode": "format", "blocks": ["p2_b2"], "note": "x"})
    f = config_prompts() / "fix_formatting.txt"
    assert f.exists() and "{note}" in f.read_text(encoding="utf-8")
    f.write_text(f.read_text(encoding="utf-8").replace("You edit the layout", "MY OWN WORDING. You edit the layout"), encoding="utf-8")
    fake_models.clear()
    ops.review(job, {"mode": "format", "blocks": ["p2_b2"], "note": "x"})
    assert "MY OWN WORDING" in fake_models[0][1]["messages"][0]["content"][0]["text"]


def test_a_broken_prompt_file_falls_back_to_the_built_in_wording(job, fake_models, capsys):
    ops.review(job, {"mode": "format", "blocks": ["p2_b2"], "note": "x"})
    (config_prompts() / "fix_formatting.txt").write_text("no placeholders here", encoding="utf-8")
    fake_models.clear()
    ops.review(job, {"mode": "format", "blocks": ["p2_b2"], "note": "x"})
    assert "THE PERSON'S INSTRUCTION" in fake_models[0][1]["messages"][0]["content"][0]["text"]
    assert "WARNING" in capsys.readouterr().err


def config_prompts():
    from ocr import config
    return config.DATA / "prompts"


# ── a small crop that Bodhan reads as nothing is read by the vision model instead ──
def test_an_empty_bodhan_answer_falls_back_to_the_vision_model(job, fake_models, monkeypatch):
    import requests
    inner = requests.post

    def post(url, *a, **kw):
        body = kw.get("json") or {}
        if "bodhan" in url:
            return inner(url, *a, **{**kw, "json": body}).__class__({"blocks": []})
        text = next((p["text"] for p in body["messages"][0]["content"] if p.get("type") == "text"), "")
        if "You are an OCR engine" in text:
            return inner(url, *a, **kw).__class__({"choices": [{"message": {"content": '{"blocks":[{"order":1,"type":"Text","text":"read by the model","bbox_norm":null}]}'}, "finish_reason": "stop"}],
                                                   "usage": {"prompt_tokens": 5, "completion_tokens": 5, "cost": 0.0}})
        return inner(url, *a, **kw)
    monkeypatch.setattr(requests, "post", post)
    res = ops.newbox(job, {"page": 1, "bbox_norm": [.1, .6, .9, .7]})
    assert res["new_block"]["formatted_text"] == "<P>read by the model</P>"
